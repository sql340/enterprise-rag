"""
企业 RAG 数据读取模块。

支持的文件类型：
- Markdown / 纯文本：.md, .markdown, .txt
- PDF：.pdf（文本层优先，扫描页可选 OCR）
- Word：.docx, .docm（.doc 通过 LibreOffice 转换）
- PowerPoint：.pptx, .pptm, .ppsx, .potx（.ppt 通过 LibreOffice 转换）
- 表格：.csv, .tsv, .xlsx, .xlsm, .xltx, .xltm, .xls, .ods
- 图片：.jpg, .jpeg, .png, .bmp, .gif, .tif, .tiff, .webp
- 音视频：.mp4, .mov, .avi, .mkv 等，以及常见音频格式

第三方依赖均按需导入，缺少依赖时会抛出带安装提示的
``OptionalDependencyError``，不会影响其它文件类型的读取。

示例::

    from data_processing.data_reader import DataReader

    reader = DataReader()
    documents = reader.read("knowledge/demo.pdf")
    for document in documents:
        print(document.metadata)
        print(document.content)

    # 接入企业 OCR / ASR 服务
    reader = DataReader(
        ocr_handler=lambda image_path: my_ocr(image_path),
        transcriber=lambda media_path: my_asr(media_path),
    )
"""

from __future__ import annotations

import csv
import io
import mimetypes
import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any, Callable, Iterable, Sequence


OCRHandler = Callable[[Path], str]
Transcriber = Callable[[Path], str]
ErrorHandler = Callable[[Path, Exception], None]

_MARKDOWN_EXTENSIONS = {".md", ".markdown"}
_TEXT_EXTENSIONS = {".txt"}
_PDF_EXTENSIONS = {".pdf"}
_WORD_EXTENSIONS = {".docx", ".docm", ".doc"}
_PPT_EXTENSIONS = {".pptx", ".pptm", ".ppsx", ".potx", ".ppt"}
_SPREADSHEET_EXTENSIONS = {
    ".csv",
    ".tsv",
    ".xlsx",
    ".xlsm",
    ".xltx",
    ".xltm",
    ".xls",
    ".ods",
}
_IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".png",
    ".bmp",
    ".gif",
    ".tif",
    ".tiff",
    ".webp",
}
_VIDEO_EXTENSIONS = {
    ".mp4",
    ".mov",
    ".avi",
    ".mkv",
    ".webm",
    ".m4v",
    ".wmv",
    ".flv",
    ".mpeg",
    ".mpg",
    ".mts",
}
_AUDIO_EXTENSIONS = {
    ".mp3",
    ".wav",
    ".m4a",
    ".aac",
    ".flac",
    ".ogg",
    ".opus",
    ".wma",
}

_KIND_BY_EXTENSION: dict[str, str] = {}
for _extensions, _kind in (
    (_MARKDOWN_EXTENSIONS, "markdown"),
    (_TEXT_EXTENSIONS, "text"),
    (_PDF_EXTENSIONS, "pdf"),
    (_WORD_EXTENSIONS, "word"),
    (_PPT_EXTENSIONS, "ppt"),
    (_SPREADSHEET_EXTENSIONS, "spreadsheet"),
    (_IMAGE_EXTENSIONS, "image"),
    (_VIDEO_EXTENSIONS, "video"),
    (_AUDIO_EXTENSIONS, "audio"),
):
    for _extension in _extensions:
        _KIND_BY_EXTENSION[_extension] = _kind

SUPPORTED_EXTENSIONS = frozenset(_KIND_BY_EXTENSION)
SUPPORTED_FILE_TYPES = frozenset(_KIND_BY_EXTENSION.values())

_FILE_TYPE_ALIASES = {
    "md": "markdown",
    "markdown": "markdown",
    "txt": "text",
    "text": "text",
    "pdf": "pdf",
    "doc": "word",
    "docx": "word",
    "word": "word",
    "ppt": "ppt",
    "pptx": "ppt",
    "powerpoint": "ppt",
    "csv": "spreadsheet",
    "tsv": "spreadsheet",
    "xls": "spreadsheet",
    "xlsx": "spreadsheet",
    "excel": "spreadsheet",
    "sheet": "spreadsheet",
    "spreadsheet": "spreadsheet",
    "image": "image",
    "img": "image",
    "video": "video",
    "audio": "audio",
}

_HEADING_STYLE_RE = re.compile(r"^(?:heading|标题)\s*([1-9])$", re.IGNORECASE)


class DataReaderError(RuntimeError):
    """数据读取过程中的基础异常。"""


class UnsupportedFileTypeError(DataReaderError, ValueError):
    """不支持的文件类型。"""


class OptionalDependencyError(DataReaderError, ImportError):
    """读取该类型文件所需的可选依赖未安装。"""


class FileReadError(DataReaderError):
    """文件存在，但解析失败。"""


@dataclass(slots=True)
class Document:
    """读取得到的一段文档内容。

    ``content`` 可直接送入后续的清洗、分块流程；``metadata`` 保留来源、
    页码、工作表、音视频时长等信息。一个文件可能产生多个 ``Document``，
    例如 PDF 的每一页、PPT 的每一页、Excel 的每一个工作表。
    """

    content: str
    metadata: dict[str, Any] = field(default_factory=dict)

    @property
    def page_content(self) -> str:
        """提供与常见 RAG 框架一致的字段别名。"""

        return self.content

    def to_dict(self) -> dict[str, Any]:
        return {"content": self.content, "metadata": dict(self.metadata)}


class DataReader:
    """按文件类型读取并转换为结构化文本。

    Args:
        encoding: 文本类文件的首选编码。读取失败时会依次尝试 UTF-8、
            GB18030、Big5 等常见编码。
        password: 加密 PDF 的密码。
        ocr_enabled: 是否对图片、扫描版 PDF 和 PPT 内图片执行 OCR。
        ocr_languages: Tesseract 语言参数，默认中英文混合。
        ocr_handler: 自定义 OCR 函数，参数为图片路径，返回识别文本。
            设置后优先于内置 Tesseract。
        transcribe_video: 是否对含音轨的音视频执行语音转写。
        transcriber: 自定义转写函数，参数为原始音视频路径，返回文本。
            设置后优先于内置 Whisper。
        whisper_model: Whisper 模型名称，默认 ``small``。
        whisper_device: faster-whisper 设备，默认 ``auto``。
        whisper_compute_type: faster-whisper 计算类型，默认 ``default``。
        language: 语音识别语言，例如 ``zh``；默认自动识别。
        include_timestamps: 内置 Whisper 转写是否添加 ``[HH:MM:SS]``。
        include_speaker_notes: 是否读取 PPT 备注。
        max_table_rows: 每个表格最多读取的行数；``None`` 表示不截断。
        office_binary: LibreOffice 可执行文件路径，用于转换旧版 Office。
        command_timeout: ffmpeg/ffprobe/LibreOffice 命令超时秒数。
        pdf_ocr_min_chars: PDF 页面文字少于该值时尝试 OCR。
    """

    def __init__(
        self,
        *,
        encoding: str = "utf-8",
        password: str | None = None,
        ocr_enabled: bool = True,
        ocr_languages: str = "chi_sim+eng",
        ocr_handler: OCRHandler | None = None,
        transcribe_video: bool = True,
        transcriber: Transcriber | None = None,
        whisper_model: str = "small",
        whisper_device: str = "auto",
        whisper_compute_type: str = "default",
        language: str | None = None,
        include_timestamps: bool = True,
        include_speaker_notes: bool = True,
        max_table_rows: int | None = None,
        office_binary: str | Path | None = None,
        command_timeout: float = 300.0,
        pdf_ocr_min_chars: int = 10,
    ) -> None:
        if max_table_rows is not None and max_table_rows <= 0:
            raise ValueError("max_table_rows must be a positive integer or None")
        if command_timeout <= 0:
            raise ValueError("command_timeout must be positive")
        if pdf_ocr_min_chars < 0:
            raise ValueError("pdf_ocr_min_chars cannot be negative")

        self.encoding = encoding
        self.password = password
        self.ocr_enabled = ocr_enabled
        self.ocr_languages = ocr_languages
        self.ocr_handler = ocr_handler
        self.transcribe_video = transcribe_video
        self.transcriber = transcriber
        self.whisper_model = whisper_model
        self.whisper_device = whisper_device
        self.whisper_compute_type = whisper_compute_type
        self.language = language
        self.include_timestamps = include_timestamps
        self.include_speaker_notes = include_speaker_notes
        self.max_table_rows = max_table_rows
        self.office_binary = Path(office_binary).expanduser() if office_binary else None
        self.command_timeout = command_timeout
        self.pdf_ocr_min_chars = pdf_ocr_min_chars

        self._whisper_engine: Any = None
        self._whisper_engine_name: str | None = None
        self._ocr_runtime_error: str | None = None

        self._handlers: dict[str, Callable[[Path], list[Document]]] = {
            "markdown": lambda path: self._read_text(path, "markdown"),
            "text": lambda path: self._read_text(path, "text"),
            "pdf": self._read_pdf,
            "word": self._read_word,
            "ppt": self._read_ppt,
            "spreadsheet": self._read_spreadsheet,
            "image": self._read_image,
            "video": self._read_media,
            "audio": self._read_media,
        }

    @property
    def supported_extensions(self) -> tuple[str, ...]:
        return tuple(sorted(SUPPORTED_EXTENSIONS))

    @property
    def supported_file_types(self) -> tuple[str, ...]:
        return tuple(sorted(SUPPORTED_FILE_TYPES))

    def supports(self, path: str | Path) -> bool:
        """判断路径扩展名是否受支持。"""

        value = str(path).strip().lower()
        extension = value if value.startswith(".") else Path(value).suffix
        return extension in SUPPORTED_EXTENSIONS

    def read(
        self,
        path: str | Path,
        *,
        file_type: str | None = None,
    ) -> list[Document]:
        """读取单个文件。

        Args:
            path: 文件路径。
            file_type: 可选的文件类型或扩展名，用于没有扩展名或扩展名
                不可靠的场景，例如 ``"pdf"``。

        Returns:
            结构化文档列表。文件可能被拆为多页、多页幻灯片或多个工作表。

        Raises:
            FileNotFoundError: 文件不存在。
            IsADirectoryError: 传入的是目录。
            UnsupportedFileTypeError: 无法识别或暂不支持该文件类型。
            OptionalDependencyError: 缺少解析该类型的可选依赖。
            FileReadError: 文件打开、解析或转换失败。
        """

        source = Path(path).expanduser()
        if not source.exists():
            raise FileNotFoundError(f"文件不存在: {source}")
        if not source.is_file():
            raise IsADirectoryError(f"需要文件路径，收到目录: {source}")

        source = source.resolve()
        kind = self._resolve_file_type(source, file_type)
        handler = self._handlers[kind]
        base_metadata = self._base_metadata(source, kind)

        try:
            documents = handler(source)
        except DataReaderError:
            raise
        except Exception as exc:
            raise FileReadError(f"读取文件失败 {source}: {exc}") from exc

        if not documents:
            documents = [Document(content="")]

        result: list[Document] = []
        for document in documents:
            content = _clean_text(document.content)
            metadata = {**base_metadata, **document.metadata}
            metadata["char_count"] = len(content)
            result.append(Document(content=content, metadata=metadata))
        return result

    def read_batch(
        self,
        paths: Iterable[str | Path],
        *,
        continue_on_error: bool = False,
        on_error: ErrorHandler | None = None,
    ) -> list[Document]:
        """读取多个文件，顺序与输入路径一致。"""

        documents: list[Document] = []
        for path in paths:
            try:
                documents.extend(self.read(path))
            except Exception as exc:
                if on_error is not None:
                    on_error(Path(path), exc)
                if not continue_on_error:
                    raise
        return documents

    def read_directory(
        self,
        directory: str | Path,
        *,
        recursive: bool = True,
        extensions: Iterable[str] | None = None,
        include_hidden: bool = False,
        continue_on_error: bool = False,
        on_error: ErrorHandler | None = None,
    ) -> list[Document]:
        """读取目录中的受支持文件。

        默认只读取已知的文件类型，并返回稳定排序后的结果，便于重复构建
        相同版本的索引。
        """

        root = Path(directory).expanduser()
        if not root.exists():
            raise FileNotFoundError(f"目录不存在: {root}")
        if not root.is_dir():
            raise NotADirectoryError(f"需要目录路径，收到文件: {root}")
        root = root.resolve()

        allowed_extensions = self._normalise_extensions(extensions)
        candidates = root.rglob("*") if recursive else root.iterdir()
        paths: list[Path] = []
        for path in candidates:
            if not path.is_file():
                continue
            relative_parts = path.relative_to(root).parts
            if not include_hidden and any(part.startswith(".") for part in relative_parts):
                continue
            if path.suffix.lower() not in allowed_extensions:
                continue
            paths.append(path)

        documents: list[Document] = []
        for path in sorted(paths, key=lambda item: str(item)):
            try:
                documents.extend(self.read(path))
            except Exception as exc:
                if on_error is not None:
                    on_error(path, exc)
                if not continue_on_error:
                    raise
        return documents

    def _normalise_extensions(self, extensions: Iterable[str] | None) -> frozenset[str]:
        if extensions is None:
            return SUPPORTED_EXTENSIONS
        if isinstance(extensions, str):
            extensions = [extensions]
        normalised: set[str] = set()
        for extension in extensions:
            value = str(extension).strip().lower()
            if not value:
                continue
            if not value.startswith("."):
                value = f".{value}"
            if value not in SUPPORTED_EXTENSIONS:
                raise UnsupportedFileTypeError(f"不支持的扩展名: {extension}")
            normalised.add(value)
        return frozenset(normalised)

    def _resolve_file_type(self, path: Path, file_type: str | None) -> str:
        if file_type is not None:
            value = str(file_type).strip().lower().lstrip(".")
            kind = _FILE_TYPE_ALIASES.get(value)
            if kind is None:
                raise UnsupportedFileTypeError(f"不支持的文件类型: {file_type}")
            return kind

        kind = _KIND_BY_EXTENSION.get(path.suffix.lower())
        if kind is None:
            supported = ", ".join(self.supported_extensions)
            raise UnsupportedFileTypeError(
                f"不支持的文件类型 {path.suffix or '(无扩展名)'}: {path}；"
                f"支持的类型: {supported}"
            )
        return kind

    def _base_metadata(self, path: Path, kind: str) -> dict[str, Any]:
        stat = path.stat()
        metadata: dict[str, Any] = {
            "source": str(path),
            "file_name": path.name,
            "extension": path.suffix.lower(),
            "file_type": kind,
            "mime_type": mimetypes.guess_type(path.name)[0]
            or "application/octet-stream",
            "size_bytes": stat.st_size,
            "modified_at": datetime.fromtimestamp(
                stat.st_mtime, tz=timezone.utc
            ).isoformat(),
        }
        if hasattr(stat, "st_birthtime"):
            metadata["created_at"] = datetime.fromtimestamp(
                stat.st_birthtime, tz=timezone.utc
            ).isoformat()
        return metadata

    # ------------------------------------------------------------------
    # Text
    # ------------------------------------------------------------------
    def _read_text(self, path: Path, kind: str) -> list[Document]:
        text = self._decode_file(path)
        metadata: dict[str, Any] = {
            "extraction_method": "native",
            "line_count": len(text.splitlines()),
        }
        return [Document(content=text, metadata=metadata)]

    def _decode_file(self, path: Path) -> str:
        return self._decode_bytes(path.read_bytes())

    def _decode_bytes(self, data: bytes) -> str:
        encodings: list[str] = []
        for encoding in (
            self.encoding,
            "utf-8-sig",
            "utf-8",
            "gb18030",
            "big5",
            "utf-16",
            "latin-1",
        ):
            if encoding and encoding.lower() not in {item.lower() for item in encodings}:
                encodings.append(encoding)

        for encoding in encodings:
            try:
                return data.decode(encoding)
            except (UnicodeDecodeError, LookupError):
                continue
        return data.decode(encodings[0], errors="replace")

    # ------------------------------------------------------------------
    # PDF
    # ------------------------------------------------------------------
    def _read_pdf(self, path: Path) -> list[Document]:
        try:
            return self._read_pdf_with_pypdf(path)
        except OptionalDependencyError:
            try:
                return self._read_pdf_with_pdfplumber(path)
            except OptionalDependencyError as exc:
                raise OptionalDependencyError(
                    "读取 PDF 需要安装 pypdf 或 pdfplumber："
                    "`pip install pypdf`。扫描版 PDF 的 OCR 还需要 "
                    "PyMuPDF（或 pdf2image + Poppler）及 Tesseract。"
                ) from exc

    def _read_pdf_with_pypdf(self, path: Path) -> list[Document]:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise OptionalDependencyError("未安装 pypdf") from exc

        reader = PdfReader(str(path))
        if reader.is_encrypted:
            password = self.password or ""
            if not reader.decrypt(password):
                raise DataReaderError("PDF 已加密，密码缺失或错误")

        document_metadata = self._pypdf_metadata(reader)
        pages = list(reader.pages)
        page_count = len(pages)
        documents: list[Document] = []
        for page_index, page in enumerate(pages):
            page_number = page_index + 1
            try:
                text = page.extract_text() or ""
                extraction_method = "pypdf"
            except Exception as exc:
                text = ""
                extraction_method = f"pypdf_error:{type(exc).__name__}"

            metadata: dict[str, Any] = {
                **document_metadata,
                "page_number": page_number,
                "page_count": page_count,
                "extraction_method": extraction_method,
            }
            if self._needs_pdf_ocr(text):
                ocr_text, ocr_error = self._ocr_pdf_page(path, page_index)
                if ocr_text:
                    text = f"{text}\n\n{ocr_text}".strip()
                    metadata["extraction_method"] = f"{extraction_method}+ocr"
                    metadata["ocr_used"] = True
                elif ocr_error:
                    metadata["ocr_error"] = ocr_error

            documents.append(Document(content=text, metadata=metadata))
        return documents

    def _read_pdf_with_pdfplumber(self, path: Path) -> list[Document]:
        try:
            import pdfplumber
        except ImportError as exc:
            raise OptionalDependencyError("未安装 pdfplumber") from exc

        try:
            pdf = pdfplumber.open(str(path), password=self.password)
        except TypeError:
            pdf = pdfplumber.open(str(path))

        with pdf:
            pages = list(pdf.pages)
            page_count = len(pages)
            documents: list[Document] = []
            for page_index, page in enumerate(pages):
                try:
                    text = page.extract_text() or ""
                    extraction_method = "pdfplumber"
                except Exception as exc:
                    text = ""
                    extraction_method = f"pdfplumber_error:{type(exc).__name__}"

                metadata: dict[str, Any] = {
                    "page_number": page_index + 1,
                    "page_count": page_count,
                    "extraction_method": extraction_method,
                }
                if self._needs_pdf_ocr(text):
                    ocr_text, ocr_error = self._ocr_pdf_page(path, page_index)
                    if ocr_text:
                        text = f"{text}\n\n{ocr_text}".strip()
                        metadata["extraction_method"] = f"{extraction_method}+ocr"
                        metadata["ocr_used"] = True
                    elif ocr_error:
                        metadata["ocr_error"] = ocr_error
                documents.append(Document(content=text, metadata=metadata))
            return documents

    @staticmethod
    def _pypdf_metadata(reader: Any) -> dict[str, Any]:
        raw = getattr(reader, "metadata", None)
        if raw is None:
            return {}
        result: dict[str, Any] = {}
        for source_key, target_key in (
            ("title", "title"),
            ("author", "author"),
            ("subject", "subject"),
            ("creator", "creator"),
            ("producer", "producer"),
        ):
            try:
                value = getattr(raw, source_key, None)
            except Exception:
                value = None
            if value:
                result[target_key] = str(value)
        return result

    def _needs_pdf_ocr(self, text: str) -> bool:
        return (
            self.ocr_enabled
            and len(_compact_text(text)) < self.pdf_ocr_min_chars
            and self._ocr_runtime_error is None
        )

    def _ocr_pdf_page(self, path: Path, page_index: int) -> tuple[str, str | None]:
        try:
            image_path, temporary_directory = self._render_pdf_page(path, page_index)
        except Exception as exc:
            error = f"PDF 页面渲染失败: {exc}"
            if isinstance(exc, OptionalDependencyError):
                self._ocr_runtime_error = str(exc)
            return "", error

        try:
            return self._extract_image_text(image_path, strict=False)
        finally:
            temporary_directory.cleanup()

    def _render_pdf_page(
        self, path: Path, page_index: int
    ) -> tuple[Path, tempfile.TemporaryDirectory[str]]:
        directory = tempfile.TemporaryDirectory(prefix="enterprise_rag_pdf_")
        output = Path(directory.name) / f"page_{page_index + 1}.png"

        fitz = None
        try:
            import pymupdf as fitz  # type: ignore[no-redef]
        except ImportError:
            try:
                import fitz  # type: ignore[no-redef]
            except ImportError:
                fitz = None

        try:
            if fitz is not None:
                document = fitz.open(str(path))
                try:
                    if getattr(document, "needs_pass", False):
                        if not document.authenticate(self.password or ""):
                            raise DataReaderError("PDF 已加密，密码缺失或错误")
                    page = document.load_page(page_index)
                    pixmap = page.get_pixmap(
                        matrix=fitz.Matrix(2.0, 2.0), alpha=False
                    )
                    pixmap.save(str(output))
                finally:
                    document.close()
                return output, directory

            try:
                from pdf2image import convert_from_path
            except ImportError as exc:
                raise OptionalDependencyError(
                    "扫描版 PDF OCR 需要安装 PyMuPDF，或安装 pdf2image 并提供 Poppler"
                ) from exc

            images = convert_from_path(
                str(path),
                first_page=page_index + 1,
                last_page=page_index + 1,
                dpi=200,
            )
            if not images:
                raise DataReaderError("未能渲染 PDF 页面")
            images[0].save(output, format="PNG")
            return output, directory
        except Exception:
            directory.cleanup()
            raise

    # ------------------------------------------------------------------
    # Word
    # ------------------------------------------------------------------
    def _read_word(self, path: Path) -> list[Document]:
        if path.suffix.lower() == ".doc":
            return self._convert_office(path, "docx", self._read_docx)
        return self._read_docx(path)

    def _read_docx(self, path: Path) -> list[Document]:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise OptionalDependencyError(
                "读取 Word 需要安装 python-docx：`pip install python-docx`"
            ) from exc

        document = DocxDocument(str(path))
        parts: list[str] = []
        paragraph_count = 0
        table_count = 0

        blocks = (
            document.iter_inner_content()
            if hasattr(document, "iter_inner_content")
            else None
        )
        if blocks is not None:
            for block in blocks:
                if hasattr(block, "text"):
                    paragraph_count += 1
                    paragraph = self._format_docx_paragraph(block)
                    if paragraph:
                        parts.append(paragraph)
                elif hasattr(block, "rows"):
                    table_count += 1
                    table = self._docx_table_to_markdown(block)
                    if table:
                        parts.append(table)
        else:
            for paragraph in document.paragraphs:
                paragraph_count += 1
                formatted = self._format_docx_paragraph(paragraph)
                if formatted:
                    parts.append(formatted)
            for table in document.tables:
                table_count += 1
                formatted_table = self._docx_table_to_markdown(table)
                if formatted_table:
                    parts.append(formatted_table)

        metadata: dict[str, Any] = {
            "extraction_method": "python-docx",
            "paragraph_count": paragraph_count,
            "table_count": table_count,
        }
        core = getattr(document, "core_properties", None)
        if core is not None:
            for attribute in ("title", "author", "subject", "keywords"):
                value = getattr(core, attribute, None)
                if value:
                    metadata[attribute] = str(value)

        return [Document(content="\n\n".join(parts), metadata=metadata)]

    @staticmethod
    def _format_docx_paragraph(paragraph: Any) -> str:
        text = (paragraph.text or "").strip()
        if not text:
            return ""

        style_name = ""
        try:
            style_name = (paragraph.style.name or "").strip()
        except Exception:
            style_name = ""

        heading_match = _HEADING_STYLE_RE.match(style_name)
        if heading_match:
            level = int(heading_match.group(1))
            return f"{'#' * level} {text}"
        if style_name.lower() in {"title", "标题"}:
            return f"# {text}"

        lowered_style = style_name.lower()
        if "list bullet" in lowered_style and not text.startswith(("- ", "* ")):
            return f"- {text}"
        if "list number" in lowered_style and not re.match(r"^\d+[.)]\s", text):
            return f"1. {text}"
        return text

    def _docx_table_to_markdown(self, table: Any) -> str:
        rows: list[list[Any]] = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        return _rows_to_markdown(rows)

    # ------------------------------------------------------------------
    # PowerPoint
    # ------------------------------------------------------------------
    def _read_ppt(self, path: Path) -> list[Document]:
        if path.suffix.lower() == ".ppt":
            return self._convert_office(path, "pptx", self._read_pptx)
        return self._read_pptx(path)

    def _read_pptx(self, path: Path) -> list[Document]:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError as exc:
            raise OptionalDependencyError(
                "读取 PowerPoint 需要安装 python-pptx：`pip install python-pptx`"
            ) from exc

        presentation = Presentation(str(path))
        documents: list[Document] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            parts: list[str] = []
            image_count = 0
            ocr_image_count = 0

            for shape in self._iter_pptx_shapes(slide.shapes, MSO_SHAPE_TYPE):
                if getattr(shape, "has_text_frame", False):
                    for paragraph in shape.text_frame.paragraphs:
                        text = (paragraph.text or "").strip()
                        if text:
                            level = max(0, int(getattr(paragraph, "level", 0) or 0))
                            parts.append(f"{'  ' * level}{text}")
                elif getattr(shape, "has_table", False):
                    table = self._pptx_table_to_markdown(shape.table)
                    if table:
                        parts.append(table)
                elif getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    image_count += 1
                    if self.ocr_enabled:
                        ocr_text, _ = self._ocr_pptx_picture(shape)
                        if ocr_text:
                            ocr_image_count += 1
                            parts.append(ocr_text)

            if self.include_speaker_notes and getattr(slide, "has_notes_slide", False):
                notes = self._pptx_notes(slide)
                if notes:
                    parts.append(f"Speaker notes:\n{notes}")

            title = ""
            try:
                if slide.shapes.title is not None:
                    title = (slide.shapes.title.text or "").strip()
            except Exception:
                title = ""

            metadata: dict[str, Any] = {
                "slide_number": slide_index,
                "slide_count": len(presentation.slides),
                "extraction_method": "python-pptx",
                "image_count": image_count,
                "ocr_image_count": ocr_image_count,
            }
            if title:
                metadata["title"] = title
            documents.append(Document(content="\n\n".join(parts), metadata=metadata))
        return documents

    @classmethod
    def _iter_pptx_shapes(
        cls, shapes: Any, shape_type_enum: Any
    ) -> Iterable[Any]:
        for shape in shapes:
            if getattr(shape, "shape_type", None) == shape_type_enum.GROUP:
                yield from cls._iter_pptx_shapes(shape.shapes, shape_type_enum)
            else:
                yield shape

    @staticmethod
    def _pptx_table_to_markdown(table: Any) -> str:
        rows = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        return _rows_to_markdown(rows)

    @staticmethod
    def _pptx_notes(slide: Any) -> str:
        try:
            text_frame = slide.notes_slide.notes_text_frame
        except Exception:
            return ""
        if text_frame is None:
            return ""
        return (text_frame.text or "").strip()

    def _ocr_pptx_picture(self, shape: Any) -> tuple[str, str | None]:
        try:
            image = shape.image
            extension = f".{image.ext}" if image.ext else ".png"
        except Exception as exc:
            return "", f"读取 PPT 图片失败: {exc}"

        with tempfile.TemporaryDirectory(prefix="enterprise_rag_ppt_") as directory:
            image_path = Path(directory) / f"image{extension}"
            image_path.write_bytes(image.blob)
            return self._extract_image_text(image_path, strict=False)

    # ------------------------------------------------------------------
    # Spreadsheet
    # ------------------------------------------------------------------
    def _read_spreadsheet(self, path: Path) -> list[Document]:
        suffix = path.suffix.lower()
        if suffix in {".csv", ".tsv"}:
            return self._read_delimited(path)
        if suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
            return self._read_xlsx(path)
        if suffix == ".xls":
            try:
                return self._read_xls(path)
            except OptionalDependencyError:
                return self._convert_office(path, "xlsx", self._read_xlsx)
        if suffix == ".ods":
            return self._convert_office(path, "xlsx", self._read_xlsx)
        raise UnsupportedFileTypeError(f"不支持的表格类型: {suffix}")

    def _read_delimited(self, path: Path) -> list[Document]:
        text = self._decode_file(path)
        if path.suffix.lower() == ".tsv":
            dialect = csv.excel_tab
        else:
            sample = text[:65536]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
            except csv.Error:
                dialect = csv.excel

        rows = list(csv.reader(io.StringIO(text), dialect))
        truncated = False
        if self.max_table_rows is not None and len(rows) > self.max_table_rows:
            rows = rows[: self.max_table_rows]
            truncated = True
        rows = _trim_rows(rows)
        metadata: dict[str, Any] = {
            "extraction_method": "csv",
            "delimiter": getattr(dialect, "delimiter", ","),
            "row_count": len(rows),
            "column_count": max((len(row) for row in rows), default=0),
            "truncated": truncated,
        }
        return [Document(content=_rows_to_markdown(rows), metadata=metadata)]

    def _read_xlsx(self, path: Path) -> list[Document]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise OptionalDependencyError(
                "读取 Excel 需要安装 openpyxl：`pip install openpyxl`"
            ) from exc

        workbook = load_workbook(
            filename=str(path), read_only=True, data_only=True, keep_links=False
        )
        try:
            documents: list[Document] = []
            for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
                rows: list[Sequence[Any]] = []
                truncated = False
                for row in worksheet.iter_rows(values_only=True):
                    if (
                        self.max_table_rows is not None
                        and len(rows) >= self.max_table_rows
                    ):
                        truncated = True
                        break
                    rows.append(row)
                rows = _trim_rows(rows)
                metadata: dict[str, Any] = {
                    "extraction_method": "openpyxl",
                    "sheet_name": worksheet.title,
                    "sheet_index": sheet_index,
                    "sheet_count": len(workbook.worksheets),
                    "sheet_state": getattr(worksheet, "sheet_state", "visible"),
                    "row_count": len(rows),
                    "column_count": max((len(row) for row in rows), default=0),
                    "truncated": truncated,
                }
                documents.append(
                    Document(content=_rows_to_markdown(rows), metadata=metadata)
                )
            return documents
        finally:
            workbook.close()

    def _read_xls(self, path: Path) -> list[Document]:
        try:
            import xlrd
        except ImportError as exc:
            raise OptionalDependencyError(
                "读取旧版 .xls 需要安装 xlrd：`pip install xlrd`"
            ) from exc

        workbook = xlrd.open_workbook(str(path))
        documents: list[Document] = []
        for sheet_index, worksheet in enumerate(workbook.sheets(), start=1):
            rows: list[list[Any]] = []
            truncated = False
            for row_index in range(worksheet.nrows):
                if (
                    self.max_table_rows is not None
                    and len(rows) >= self.max_table_rows
                ):
                    truncated = True
                    break
                row: list[Any] = []
                for column_index in range(worksheet.ncols):
                    cell = worksheet.cell(row_index, column_index)
                    value: Any = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            value = xlrd.xldate_as_datetime(
                                cell.value, workbook.datemode
                            )
                        except Exception:
                            pass
                    row.append(value)
                rows.append(row)
            rows = _trim_rows(rows)
            metadata: dict[str, Any] = {
                "extraction_method": "xlrd",
                "sheet_name": worksheet.name,
                "sheet_index": sheet_index,
                "sheet_count": len(workbook.sheets()),
                "row_count": len(rows),
                "column_count": max((len(row) for row in rows), default=0),
                "truncated": truncated,
            }
            documents.append(Document(content=_rows_to_markdown(rows), metadata=metadata))
        return documents

    # ------------------------------------------------------------------
    # Image
    # ------------------------------------------------------------------
    def _read_image(self, path: Path) -> list[Document]:
        try:
            from PIL import Image, ImageOps
        except ImportError as exc:
            raise OptionalDependencyError(
                "读取图片需要安装 Pillow：`pip install pillow`"
            ) from exc

        with Image.open(path) as image:
            oriented = ImageOps.exif_transpose(image)
            width, height = oriented.size
            image_format = image.format or path.suffix.lstrip(".").upper()
            mode = image.mode

        text = ""
        extraction_method = "metadata"
        ocr_error: str | None = None
        if self.ocr_enabled:
            text, ocr_error = self._extract_image_text(path, strict=True)
            extraction_method = "ocr" if text else "ocr_empty"

        metadata: dict[str, Any] = {
            "extraction_method": extraction_method,
            "image_format": image_format,
            "width": width,
            "height": height,
            "color_mode": mode,
        }
        if ocr_error:
            metadata["ocr_error"] = ocr_error
        return [Document(content=text, metadata=metadata)]

    def _extract_image_text(self, path: Path, *, strict: bool) -> tuple[str, str | None]:
        if not self.ocr_enabled:
            return "", None
        try:
            if self.ocr_handler is not None:
                text = self.ocr_handler(path)
                return _clean_text(str(text or "")), None
            text = self._pytesseract_ocr(path)
            return _clean_text(text), None
        except OptionalDependencyError:
            if strict:
                raise
            return "", self._ocr_runtime_error or "OCR 依赖不可用"
        except Exception as exc:
            if strict:
                raise FileReadError(f"OCR 识别失败 {path}: {exc}") from exc
            return "", str(exc)

    def _pytesseract_ocr(self, path: Path) -> str:
        try:
            import pytesseract
            from PIL import Image, ImageOps
        except ImportError as exc:
            self._ocr_runtime_error = (
                "读取图片内容需要安装 Pillow 和 pytesseract，并安装 Tesseract OCR："
                "`pip install pillow pytesseract`"
            )
            raise OptionalDependencyError(self._ocr_runtime_error) from exc

        with Image.open(path) as image:
            image = ImageOps.exif_transpose(image)
            if image.mode not in {"RGB", "L"}:
                image = image.convert("RGB")
            try:
                return pytesseract.image_to_string(image, lang=self.ocr_languages)
            except pytesseract.TesseractNotFoundError as exc:
                self._ocr_runtime_error = (
                    "未找到 Tesseract 可执行文件，请安装 Tesseract OCR，"
                    "或通过 ocr_handler 注入自定义 OCR 服务"
                )
                raise OptionalDependencyError(self._ocr_runtime_error) from exc
            except pytesseract.TesseractError as exc:
                raise DataReaderError(f"Tesseract OCR 执行失败: {exc}") from exc

    # ------------------------------------------------------------------
    # Audio / video
    # ------------------------------------------------------------------
    def _read_media(self, path: Path) -> list[Document]:
        metadata, probe_error = self._probe_media(path)
        metadata["extraction_method"] = "metadata"
        if probe_error:
            metadata["probe_error"] = probe_error

        has_audio = metadata.get("has_audio")
        should_transcribe = self.transcribe_video and has_audio is not False
        text = ""
        if should_transcribe:
            text, transcription_method = self._transcribe_media(path)
            metadata["extraction_method"] = transcription_method
            metadata["ocr_used"] = False
        elif self.transcribe_video:
            metadata["transcription_skipped"] = "media has no audio stream"

        return [Document(content=text, metadata=metadata)]

    def _probe_media(self, path: Path) -> tuple[dict[str, Any], str | None]:
        ffprobe = self._find_binary("ffprobe")
        if ffprobe is None:
            return {}, "未找到 ffprobe，仅返回基础文件元数据"

        command = [
            ffprobe,
            "-v",
            "error",
            "-show_streams",
            "-show_format",
            "-of",
            "json",
            str(path),
        ]
        try:
            result = subprocess.run(
                command,
                check=False,
                capture_output=True,
                text=True,
                timeout=self.command_timeout,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            return {}, f"ffprobe 执行失败: {exc}"
        if result.returncode != 0:
            return {}, (result.stderr or "ffprobe 执行失败").strip()

        try:
            import json

            data = json.loads(result.stdout or "{}")
        except ValueError as exc:
            return {}, f"ffprobe 返回内容无法解析: {exc}"

        streams = data.get("streams") or []
        format_info = data.get("format") or {}
        video_stream = next(
            (stream for stream in streams if stream.get("codec_type") == "video"),
            None,
        )
        audio_stream = next(
            (stream for stream in streams if stream.get("codec_type") == "audio"),
            None,
        )

        metadata: dict[str, Any] = {
            "has_video": video_stream is not None,
            "has_audio": audio_stream is not None,
            "format_name": format_info.get("format_name"),
            "duration_seconds": _to_float(
                format_info.get("duration")
                or (video_stream or audio_stream or {}).get("duration")
            ),
            "bit_rate": _to_int(format_info.get("bit_rate")),
        }
        if video_stream is not None:
            metadata.update(
                {
                    "video_codec": video_stream.get("codec_name"),
                    "width": _to_int(video_stream.get("width")),
                    "height": _to_int(video_stream.get("height")),
                    "frame_rate": video_stream.get("avg_frame_rate"),
                }
            )
        if audio_stream is not None:
            metadata.update(
                {
                    "audio_codec": audio_stream.get("codec_name"),
                    "sample_rate": _to_int(audio_stream.get("sample_rate")),
                    "channels": _to_int(audio_stream.get("channels")),
                }
            )
        tags = format_info.get("tags") or {}
        if tags.get("title"):
            metadata["title"] = str(tags["title"])

        return _drop_none(metadata), None

    def _transcribe_media(self, path: Path) -> tuple[str, str]:
        if self.transcriber is not None:
            text = self.transcriber(path)
            self._validate_transcription(text)
            return _clean_text(str(text or "")), "custom_transcriber"

        with tempfile.TemporaryDirectory(prefix="enterprise_rag_media_") as directory:
            audio_path = self._extract_audio(path, Path(directory))
            text, engine = self._transcribe_with_whisper(audio_path)
            self._validate_transcription(text)
            return text, engine

    @staticmethod
    def _validate_transcription(text: str) -> None:
        if not isinstance(text, str):
            raise DataReaderError("转写器必须返回字符串")

    def _extract_audio(self, path: Path, directory: Path) -> Path:
        ffmpeg = self._find_binary("ffmpeg")
        if ffmpeg is None:
            raise OptionalDependencyError(
                "音视频转写需要安装 ffmpeg，或通过 transcriber 注入自定义 ASR 服务"
            )
        output = directory / "audio.wav"
        command = [
            ffmpeg,
            "-y",
            "-loglevel",
            "error",
            "-i",
            str(path),
            "-vn",
            "-ac",
            "1",
            "-ar",
            "16000",
            "-f",
            "wav",
            str(output),
        ]
        try:
            result = subprocess.run(
                command,
                check=False,
                capture_output=True,
                text=True,
                timeout=self.command_timeout,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            raise DataReaderError(f"提取音频失败: {exc}") from exc
        if result.returncode != 0 or not output.exists():
            raise DataReaderError(
                f"提取音频失败: {(result.stderr or 'ffmpeg 执行失败').strip()}"
            )
        return output

    def _transcribe_with_whisper(self, audio_path: Path) -> tuple[str, str]:
        if self._whisper_engine is None:
            try:
                from faster_whisper import WhisperModel

                self._whisper_engine = WhisperModel(
                    self.whisper_model,
                    device=self.whisper_device,
                    compute_type=self.whisper_compute_type,
                )
                self._whisper_engine_name = "faster-whisper"
            except ImportError:
                try:
                    import whisper

                    self._whisper_engine = whisper.load_model(self.whisper_model)
                    self._whisper_engine_name = "openai-whisper"
                except ImportError as exc:
                    raise OptionalDependencyError(
                        "音视频转写需要安装 faster-whisper 或 openai-whisper："
                        "`pip install faster-whisper`"
                    ) from exc

        if self._whisper_engine_name == "faster-whisper":
            segments, _ = self._whisper_engine.transcribe(
                str(audio_path),
                language=self.language,
                beam_size=5,
                vad_filter=True,
            )
            lines: list[str] = []
            for segment in segments:
                text = (segment.text or "").strip()
                if not text:
                    continue
                if self.include_timestamps:
                    text = f"[{_format_timestamp(segment.start)}] {text}"
                lines.append(text)
            return _clean_text("\n".join(lines)), "faster-whisper"

        result = self._whisper_engine.transcribe(
            str(audio_path), language=self.language
        )
        lines = []
        for segment in result.get("segments") or []:
            text = str(segment.get("text") or "").strip()
            if not text:
                continue
            if self.include_timestamps:
                start = _to_float(segment.get("start")) or 0.0
                text = f"[{_format_timestamp(start)}] {text}"
            lines.append(text)
        if not lines and result.get("text"):
            lines.append(str(result["text"]).strip())
        return _clean_text("\n".join(lines)), "openai-whisper"

    def _find_binary(self, name: str) -> str | None:
        return shutil.which(name)

    # ------------------------------------------------------------------
    # Legacy Office conversion
    # ------------------------------------------------------------------
    def _convert_office(
        self,
        path: Path,
        target_extension: str,
        handler: Callable[[Path], list[Document]],
    ) -> list[Document]:
        office = self._find_office_binary()
        if office is None:
            raise OptionalDependencyError(
                f"读取 {path.suffix} 需要安装 LibreOffice 进行格式转换，"
                "或先手动转换为新版 Office 格式"
            )

        with tempfile.TemporaryDirectory(prefix="enterprise_rag_office_") as directory:
            command = [
                office,
                "--headless",
                "--norestore",
                "--convert-to",
                target_extension,
                "--outdir",
                directory,
                str(path),
            ]
            try:
                result = subprocess.run(
                    command,
                    check=False,
                    capture_output=True,
                    text=True,
                    timeout=self.command_timeout,
                )
            except (OSError, subprocess.TimeoutExpired) as exc:
                raise DataReaderError(f"LibreOffice 转换失败: {exc}") from exc
            if result.returncode != 0:
                raise DataReaderError(
                    "LibreOffice 转换失败: "
                    f"{(result.stderr or result.stdout or 'unknown error').strip()}"
                )

            converted_files = sorted(
                Path(directory).glob(f"*.{target_extension}")
            )
            if not converted_files:
                raise DataReaderError(
                    f"LibreOffice 未生成 .{target_extension} 文件: {path}"
                )
            return handler(converted_files[0])

    def _find_office_binary(self) -> str | None:
        if self.office_binary is not None:
            if self.office_binary.is_file():
                return str(self.office_binary)
            found = shutil.which(str(self.office_binary))
            if found:
                return found
            raise OptionalDependencyError(
                f"指定的 LibreOffice 路径不存在: {self.office_binary}"
            )

        for candidate in ("soffice", "libreoffice"):
            found = shutil.which(candidate)
            if found:
                return found
        mac_path = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
        if mac_path.is_file():
            return str(mac_path)
        return None


def read_file(
    path: str | Path,
    *,
    file_type: str | None = None,
    **reader_options: Any,
) -> list[Document]:
    """使用默认配置或传入的 ``DataReader`` 参数读取单个文件。"""

    return DataReader(**reader_options).read(path, file_type=file_type)


def read_directory(
    directory: str | Path,
    *,
    recursive: bool = True,
    extensions: Iterable[str] | None = None,
    include_hidden: bool = False,
    continue_on_error: bool = False,
    on_error: ErrorHandler | None = None,
    **reader_options: Any,
) -> list[Document]:
    """使用默认配置或传入的 ``DataReader`` 参数读取目录。"""

    reader = DataReader(**reader_options)
    return reader.read_directory(
        directory,
        recursive=recursive,
        extensions=extensions,
        include_hidden=include_hidden,
        continue_on_error=continue_on_error,
        on_error=on_error,
    )


def _rows_to_markdown(rows: Sequence[Sequence[Any]]) -> str:
    if not rows:
        return ""

    width = max((len(row) for row in rows), default=0)
    if width == 0:
        return ""
    normalised = [
        [_cell_to_text(row[index]) if index < len(row) else "" for index in range(width)]
        for row in rows
    ]
    header = normalised[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in normalised[1:])
    return "\n".join(lines)


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, (date, time)):
        return value.isoformat()
    text = str(value)
    return (
        text.replace("\r\n", "\n")
        .replace("\r", "\n")
        .replace("|", "\\|")
        .replace("\n", "<br>")
        .strip()
    )


def _trim_rows(rows: Sequence[Sequence[Any]]) -> list[list[Any]]:
    result = [list(row) for row in rows]
    while result and all(_cell_is_empty(value) for value in result[-1]):
        result.pop()
    if not result:
        return []

    max_columns = 0
    for row in result:
        for index, value in enumerate(row):
            if not _cell_is_empty(value):
                max_columns = max(max_columns, index + 1)
    if max_columns == 0:
        return []
    return [row[:max_columns] for row in result]


def _cell_is_empty(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _clean_text(text: str) -> str:
    return (
        str(text or "")
        .replace("\x00", "")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
        .strip()
    )


def _compact_text(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def _format_timestamp(seconds: float) -> str:
    total_seconds = max(0, int(seconds or 0))
    hours, remainder = divmod(total_seconds, 3600)
    minutes, seconds = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}"


def _to_float(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _to_int(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _drop_none(mapping: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in mapping.items() if value is not None}


__all__ = [
    "DataReader",
    "DataReaderError",
    "Document",
    "FileReadError",
    "OptionalDependencyError",
    "SUPPORTED_EXTENSIONS",
    "SUPPORTED_FILE_TYPES",
    "UnsupportedFileTypeError",
    "read_directory",
    "read_file",
]
